#!/usr/bin/env python3
"""Formula test: build both workbooks, recalculate them in LibreOffice (headless), and compare every
total against the same sums done in Python from the input cells.   python3 test_xlsx.py"""
import subprocess, sys, tempfile
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).parent))
import build_xlsx as B

tmp = Path(tempfile.mkdtemp())
subprocess.run([sys.executable, str(Path(__file__).parent / "build_xlsx.py"), str(tmp)], check=True, capture_output=True)
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "r"),
                *map(str, tmp.glob("*.xlsx"))], check=True, capture_output=True)
ok = True
def eq(label, got, want):
    global ok
    good = (abs((got or 0) - want) < 0.005) if isinstance(want, (int, float)) else got == want
    if not good:
        ok = False; print(f"FAIL {label}: got {got!r}, want {want!r}")

for name in ("Christmas-Gift-Budget-Tracker.xlsx", "Christmas-Gift-Budget-Tracker-Example.xlsx"):
    src = load_workbook(tmp / name)                 # inputs
    wb = load_workbook(tmp / "r" / name, data_only=True)   # recalculated values
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") or c.value in ("Err:502", "Err:504"):
                    ok = False; print("FAIL error value", name, ws.title, c.coordinate, c.value)
    g = [[src["Gift Tracker"].cell(r, c).value for c in range(1, 8)] for r in range(B.G0, B.G1 + 1)]
    g = [x for x in g if any(v is not None for v in x)]
    ppl = [[src["People"].cell(r, c).value for c in range(1, 4)] for r in range(B.P0, B.P1 + 1)]
    ppl = [x for x in ppl if x[0]]
    ex = [[src["Expense Log"].cell(r, c).value for c in range(1, 5)] for r in range(B.E0, B.E1 + 1)]
    ex = [x for x in ex if x[3] is not None]
    cats = [(src["Budget"].cell(r, 1).value, src["Budget"].cell(r, 2).value or 0) for r in range(B.B0, B.B1 + 1)]
    P, Bu, D = wb["People"], wb["Budget"], wb["Dashboard"]
    for i, (n, grp, bud) in enumerate(ppl):
        r = B.P0 + i
        mine = [x for x in g if x[0] == n]
        planned = sum(x[3] or 0 for x in mine); spent = sum(x[3] or 0 for x in mine if x[4] == "Yes")
        eq(f"{n} planned", P.cell(r, 4).value, planned); eq(f"{n} spent", P.cell(r, 5).value, spent)
        eq(f"{n} left", P.cell(r, 6).value, bud - spent)
        for col, k in ((7, None), (8, 4), (9, 5), (10, 6)):
            eq(f"{n} count{col}", P.cell(r, col).value, len(mine) if k is None else sum(1 for x in mine if x[k] == "Yes"))
        G, H, I, J = len(mine), *(sum(1 for x in mine if x[k] == "Yes") for k in (4, 5, 6))
        want = ("No gift yet" if G == 0 else "Over budget" if spent > bud else "Plan over budget" if planned > bud and J < G
                else "All given ✓" if J == G else "Shopping" if H < G else "To wrap" if I < G else "Ready to give")
        eq(f"{n} status", P.cell(r, 11).value, want)
    gift_budget = sum(x[2] or 0 for x in ppl)
    gift_spent = sum(x[3] or 0 for x in g if x[4] == "Yes")
    eq("budget gifts spent", Bu.cell(B.B_GIFT, 3).value, gift_spent)
    for i, (c, b) in enumerate(cats):
        if c:
            eq(f"budget {c}", Bu.cell(B.B0 + i, 3).value, sum(x[3] for x in ex if x[2] == c))
    total_b = gift_budget + sum(b for c, b in cats); total_s = gift_spent + sum(x[3] for x in ex)
    eq("total budget", D["B7"].value, total_b); eq("total spent", D["B8"].value, total_s)
    eq("left", D["B9"].value, total_b - total_s)
    eq("gifts", D["E8"].value, sum(1 for x in g if x[1])); eq("bought", D["E9"].value, sum(1 for x in g if x[4] == "Yes"))
    eq("uncategorized", D["B20"].value, 0); eq("unknown names", D["B21"].value, 0)
    today = date.today(); xmas = date(today.year, 12, 25)
    if today > xmas: xmas = date(today.year + 1, 12, 25)
    eq("days", D["A4"].value, (xmas - today).days)
    S = wb["Savings"]; goal = src["Savings"]["B4"].value
    if goal:
        st = src["Savings"]["B5"].value.date() if isinstance(src["Savings"]["B5"].value, datetime) else src["Savings"]["B5"].value
        import math
        weeks = max(1, math.ceil((date(st.year, 12, 25) - st).days / 7))
        eq("weeks", S["B7"].value, weeks); eq("per week", S["B8"].value, goal / weeks)
        saved = sum(src["Savings"].cell(r, 4).value or 0 for r in range(B.S0, B.S1 + 1))
        eq("saved", S["D4"].value, saved); eq("dash saved", D["B15"].value, saved)
        eq("last week row", S.cell(B.S0 + weeks - 1, 1).value, weeks); eq("no extra week", S.cell(B.S0 + weeks, 1).value, None)
    print(name, "checked:", len(ppl), "people,", len(g), "gifts,", len(ex), "expenses")
print("ALL OK" if ok else "FAILURES"); sys.exit(0 if ok else 1)
