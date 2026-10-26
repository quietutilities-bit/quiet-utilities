#!/usr/bin/env python3
"""Turn the recalculated example workbook into listing/data.js (cell text as the sheet shows it),
so the listing mockups show the real file's numbers.   python3 gen_listing_data.py <recalculated-example.xlsx>"""
import json, sys
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
import build_xlsx as B

wb = load_workbook(sys.argv[1], data_only=True)
def fmt(v, kind=""):
    if v is None or v == "": return ""
    if kind == "$": return ("-" if v < 0 else "") + "${:,.2f}".format(abs(v))
    if kind == "%": return f"{round(v * 100)}%"
    if kind == "d": return v.strftime("%b %-d, %Y") if isinstance(v, datetime) else str(v)
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v)
def rows(ws, r0, r1, kinds):
    return [[fmt(ws.cell(r, c + 1).value, k) for c, k in enumerate(kinds)] for r in range(r0, r1 + 1)]
D = wb["Dashboard"]
data = {
  "people": rows(wb["People"], B.P0, B.P0 + 13, ["", "", "$", "$", "$", "$", "", "", "", "", ""]),
  "peopleTotal": rows(wb["People"], B.P1 + 2, B.P1 + 2, ["", "", "$", "$", "$", "$", "", "", "", "", ""])[0],
  "gifts": rows(wb["Gift Tracker"], B.G0, B.G0 + 25, ["", "", "", "$", "", "", "", ""]),
  "budget": rows(wb["Budget"], B.B_GIFT, B.B_TOT, ["", "$", "$", "$", "%", ""]),
  "expenses": rows(wb["Expense Log"], B.E0, B.E0 + 13, ["d", "", "", "$", "", "$"]),
  "savingsTop": {k: fmt(wb["Savings"][k].value, t) for k, t in
                 [("B4", "$"), ("B5", "d"), ("B6", "d"), ("B7", ""), ("B8", "$"), ("D4", "$"), ("D5", "$"), ("D6", "%"), ("E6", "")]},
  "savings": rows(wb["Savings"], B.S0, B.S0 + 15, ["", "d", "$", "$", "$", ""]),
  "dash": {k: fmt(D[k].value, t) for k, t in
           [("B7", "$"), ("B8", "$"), ("B9", "$"), ("B10", "%"), ("A11", ""), ("E7", ""), ("E8", ""), ("E9", ""), ("F9", ""),
            ("E10", ""), ("F10", ""), ("E11", ""), ("F11", ""), ("E12", ""), ("B14", "$"), ("B15", "$"), ("A16", ""), ("B20", "$"), ("B21", "")]},
  "dashCats": [[fmt(D.cell(r, c).value, k) for c, k in ((4, ""), (5, "$"), (6, ""))] for r in range(16, 29)],
}
data["dash"]["A4"] = "60"   # fixed for the photos (the file itself counts down from today)
Path(__file__).parent.joinpath("listing", "data.js").write_text("window.DATA = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")
print("listing/data.js written")
